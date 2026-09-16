# services/reports_email_service.py
"""
Geração dos e-mails matinais (Report Ibéria/Brasil, CAB, P1 Semanal) via
automação COM do Outlook local — porta de uma ferramenta standalone
(app.py + gcc_tools.html, Flask + pywin32) que existia à parte deste
projeto. Unificada aqui pra correr no mesmo backend FastAPI que já serve
o dashboard, em vez de um segundo processo Flask.

Só funciona numa máquina Windows com Outlook instalado (win32com.client
faz Dispatch("Outlook.Application") no processo local) — consistente com
o resto do projeto, que já assume Windows (xlrd/openpyxl, caminhos).

Contactos (To/CC) e as imagens de cabeçalho/rodapé do CAB são lidos de
backend/data/reports/ — ver README.md nessa pasta. Ambos têm fallback
gracioso (endereços fixos / cabeçalho só texto) se os ficheiros reais
ainda não tiverem sido colocados lá (não estão no git — contêm dados
reais, ver .gitignore).
"""
import html
import logging
from datetime import datetime
from pathlib import Path

import openpyxl
import pythoncom
import win32com.client

logger = logging.getLogger("reports_email")


def _esc(value) -> str:
    """
    Escapa um valor antes de o meter no HTML do e-mail (& < > " ').
    RESOLVIDO (bug real, revisão de segurança 2026-09): nenhuma destas
    funções escapava nada antes disto — células de tabela vêm de exports
    do ServiceNow (texto livre em "Short description"/"Motivo"/etc.) ou
    de entradas manuais no CAB (ManualEntryForm.jsx no frontend, sem
    validação nenhuma), então um "<"/">" nesses campos já bastava para
    partir a tabela do e-mail, e um valor malicioso podia injetar
    HTML/links no rascunho gerado no Outlook antes de alguém o rever e
    enviar. `mail.HTMLBody` é a única função de e-mail deste ficheiro
    afetada — `ponto_situacao_service.gerar_email_ponto_situacao` usa
    `mail.Body` (texto simples, sem interpretação de HTML), não precisa
    disto.
    """
    return html.escape(str(value)) if value is not None else ""

BASE_DIR = Path(__file__).resolve().parent.parent
REPORTS_DATA_DIR = BASE_DIR / "data" / "reports"
IMGS_DIR = REPORTS_DATA_DIR / "imgs"
CONTACTOS_PATH = REPORTS_DATA_DIR / "Contactos.xlsx"


def _imagem_para_base64(nome_ficheiro: str) -> str:
    """Lê uma imagem de data/reports/imgs/ e devolve string base64 (ou "" se não existir)."""
    import base64

    caminho = IMGS_DIR / nome_ficheiro
    if not caminho.is_file():
        return ""
    return base64.b64encode(caminho.read_bytes()).decode("ascii")


def _ler_contactos(sheet_name: str, fallback_to: str = "", fallback_cc: str = "") -> tuple[str, str]:
    """
    Lê To (A2) e CC (B2) da sheet indicada em data/reports/Contactos.xlsx.

    Sheets usadas:
        'CAB'             -> Report CAB Diário / FDS
        'P1 SEMANA'       -> Report P1 Semanal
        'CALL MATINAL PT' -> Report Ibéria (diário e FDS)
        'CALL MATINAL BR' -> Report Brasil (diário e FDS)

    Devolve (str_to, str_cc). Se não conseguir ler usa os fallbacks.
    """
    if not CONTACTOS_PATH.is_file():
        return fallback_to, fallback_cc

    try:
        wb = openpyxl.load_workbook(CONTACTOS_PATH, read_only=True, data_only=True)
        found = next(
            (s for s in wb.sheetnames if s.strip().upper() == sheet_name.strip().upper()),
            None,
        )
        if found is None:
            wb.close()
            return fallback_to, fallback_cc
        ws = wb[found]
        str_to = str(ws.cell(row=2, column=1).value or "").strip()
        str_cc = str(ws.cell(row=2, column=2).value or "").strip()
        wb.close()
    except Exception:
        logger.exception("Falha a ler Contactos.xlsx (sheet '%s')", sheet_name)
        return fallback_to, fallback_cc

    return str_to or fallback_to, str_cc or fallback_cc


def _nome_mes(m: int) -> str:
    meses = ["", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
             "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
    return meses[m] if 1 <= m <= 12 else ""


# ------------------------------------------------------------------------------
# MONTAGEM DE HTML
# ------------------------------------------------------------------------------

def montar_corpo_html_p1_semanal(linhas: list[dict]) -> str:
    total = len(linhas)

    if total == 0:
        return "Bom dia,<br><br>Não existiram P1's para ambas as geografias (Ibéria e Brasil) na semana passada."

    texto_intro = (
        "Bom dia,<br><br>Na semana passada tivemos o seguinte P1.<br><br>"
        if total == 1 else
        "Bom dia,<br><br>Na semana passada tivemos os seguintes P1.<br><br>"
    )

    cabecalhos = list(linhas[0].keys())
    coluna_number = next((c for c in cabecalhos if c.strip().lower() == "number"), None)

    html = texto_intro
    html += (
        "<table style='border-collapse:collapse; font-family:\"EDP Preon\", Calibri, "
        "Arial, sans-serif; font-size:13px; width:100%; "
        "box-shadow: 0 2px 4px rgba(0,0,0,0.1);'>"
    )
    html += "<tr style='background-color:#1F3864; color:#FFFFFF;'>"
    for cab in cabecalhos:
        html += f"<th style='padding:10px 12px; text-align:left; border:1px solid #14274e;'>{_esc(cab)}</th>"
    html += "</tr>"

    for item in linhas:
        html += "<tr style='background-color:#FFFFFF;'>"
        for cab in cabecalhos:
            val = item.get(cab, "")
            estilo = "padding:8px 12px; border:1px solid #BDD7EE; color:#1F3864;"
            if cab == coluna_number:
                estilo += " font-weight:bold;"
            html += f"<td style='{estilo}'>{_esc(val)}</td>"
        html += "</tr>"

    html += "</table>"
    return html


def montar_corpo_html_cab(tipo: str, texto_inicio: str, texto_fim: str, changes: list[dict]) -> str:
    """
    Header: imgs/img.png (578x96) sobre fundo #37B4C3
    Corpo : texto + cada change
    Footer: imgs/img2.png (578x56) sobre fundo #37B4C3
    """
    b64_img1 = _imagem_para_base64("img.png")
    b64_img2 = _imagem_para_base64("img2.png")

    header_img = (
        f"<img src='data:image/png;base64,{b64_img1}' "
        f"alt='Centro de Comando - informação' width='578' height='96' "
        f"style='display:block;border:0;outline:none;width:578px;height:96px;margin:0 auto;' border='0'>"
        if b64_img1 else
        "<p style='margin:0;font-size:20px;font-weight:bold;color:#ffffff;letter-spacing:1px;'>"
        "Global Command Center EDP</p>"
        "<p style='margin:4px 0 0 0;font-size:13px;color:#e0f7fa;'>Change Advisory Board</p>"
    )

    footer_img = (
        f"<img src='data:image/png;base64,{b64_img2}' "
        f"alt='EDP' width='578' height='56' "
        f"style='display:block;border:0;outline:none;width:578px;height:56px;margin:0 auto;' border='0'>"
        if b64_img2 else
        "<p style='margin:0;font-size:12px;color:#e0f7fa;'>"
        "EDP - Global Command Center</p>"
    )

    periodo_txt = (
        f"A execução da lista de CHG abaixo ocorrerá entre as 18:00 Horas de "
        f"{_esc(texto_inicio)} e as 07:00 de {_esc(texto_fim)}."
    )

    if not changes:
        changes_html = "<p style='font-size:14px;color:#333;'>Sem changes aprovados para o período.</p>"
    else:
        changes_html = ""
        for row in changes:
            # Colunas detectadas por linha: changes importados (Planned start/end date)
            # e changes manuais (Unavailability Start/End Date) usam nomes diferentes.
            cols = list(row.keys())
            col_num = next((c for c in cols if c.strip().lower() == "number"), None)
            col_ci = next((c for c in cols if any(k in c.lower() for k in ("item", "cmdb", " ci"))), None)
            col_desc = next((c for c in cols if any(k in c.lower() for k in ("description", "desc"))), None)
            col_start = next((c for c in cols if any(k in c.lower() for k in ("start", "begin"))), None)
            col_end = next((c for c in cols if any(k in c.lower() for k in ("end date", "end"))), None)

            numero = str(row.get(col_num, "") if col_num else "")
            ci = str(row.get(col_ci, "") if col_ci else "")
            desc = str(row.get(col_desc, "") if col_desc else "")
            inicio = str(row.get(col_start, "") if col_start else "")
            fim = str(row.get(col_end, "") if col_end else "")
            cabecalho = " - ".join(filter(None, [numero, ci, desc]))
            changes_html += (
                f"<div style='margin:12px 0; padding:10px 0; border-bottom:1px solid #e0e0e0;'>"
                f"<p style='margin:0 0 4px 0; font-size:16px; font-weight:bold; color:#222;'>{_esc(cabecalho)}</p>"
                f"<p style='margin:0; font-size:14px; color:#555;'>"
                f"Unavailability Start Date: {_esc(inicio)}<br>"
                f"Unavailability End Date: {_esc(fim)}"
                f"</p></div>"
            )

    return (
        "<!DOCTYPE html>"
        "<html xmlns:v='urn:schemas-microsoft-com:vml'"
        "      xmlns:o='urn:schemas-microsoft-com:office:office'"
        "      xmlns='http://www.w3.org/TR/REC-html40'>"
        "<head>"
        "<meta charset='UTF-8'>"
        "<meta http-equiv='X-UA-Compatible' content='IE=edge'>"
        "<meta name='color-scheme' content='light only'>"
        "<meta name='supported-color-schemes' content='light'>"
        "<meta name='format-detection' content='telephone=no,date=no,address=no,email=no,url=no'>"
        "<!--[if gte mso 9]><xml><o:OfficeDocumentSettings><o:AllowPNG/></o:OfficeDocumentSettings></xml><![endif]-->"
        "<style>"
        ":root{color-scheme:light only;}"
        "body{background-color:#ffffff !important;color:#222222 !important;font-family:Arial,sans-serif;margin:0;padding:0;}"
        "[data-ogsc] .hdr-bg,[data-ogsb] .hdr-bg{background-color:#37B4C3 !important;}"
        "[data-ogsc] .bdy-bg,[data-ogsb] .bdy-bg{background-color:#ffffff !important;}"
        "[data-ogsc] .t-dark,[data-ogsb] .t-dark{color:#222222 !important;-webkit-text-fill-color:#222222 !important;}"
        "@media(prefers-color-scheme:dark){body,table,td{background-color:#ffffff !important;color:#222222 !important;}"
        ".hdr-bg{background-color:#37B4C3 !important;}.t-dark{color:#222222 !important;}}"
        "[data-ogsc] *,[data-ogsb] *{color-scheme:light only !important;}"
        ".ExternalClass *{background-color:#ffffff !important;color:#222222 !important;}"
        "*{forced-color-adjust:none !important;}"
        "</style>"
        "<!--[if mso]><style type='text/css'>"
        ".hdr-bg{background:#37B4C3 !important;}.bdy-bg{background:#ffffff !important;}.t-dark{color:#222222 !important;}"
        "</style><![endif]-->"
        "</head>"
        "<body bgcolor='#ffffff' text='#222222' class='bdy-bg' data-ogsc='#ffffff' data-ogsb='#ffffff'"
        "  style='margin:0;padding:0;font-family:Arial,sans-serif;"
        "         background-color:#ffffff !important;color:#222222 !important;'>"

        # HEADER
        "<table width='100%' cellpadding='0' cellspacing='0' border='0'"
        "       bgcolor='#37B4C3' class='hdr-bg'"
        "       style='background-color:#37B4C3 !important;'>"
        "<tr><td align='center' valign='middle' bgcolor='#37B4C3' class='hdr-bg'"
        "    style='background-color:#37B4C3 !important;padding:0;margin:0;height:140px;'>"
        f"{header_img}"
        "</td></tr></table>"

        # CORPO
        "<table width='100%' cellpadding='0' cellspacing='0' border='0'"
        "       bgcolor='#ffffff' class='bdy-bg'"
        "       style='background-color:#ffffff !important;'>"
        "<tr><td bgcolor='#ffffff' class='bdy-bg'"
        "    style='padding:30px 40px;background-color:#ffffff !important;'>"
        "<p class='t-dark' style='font-size:16px;color:#222222 !important;mso-color-alt:#222222;margin:0 0 16px 0;'>Boa tarde,</p>"
        "<p class='t-dark' style='font-size:16px;color:#222222 !important;mso-color-alt:#222222;margin:0 0 16px 0;'>Vimos por este meio apresentar o plano de ações para hoje.</p>"
        f"<p class='t-dark' style='font-size:16px;color:#222222 !important;mso-color-alt:#222222;margin:0 0 16px 0;'>{periodo_txt}</p>"
        "<p class='t-dark' style='font-size:16px;font-weight:bold;color:#222222 !important;mso-color-alt:#222222;margin:25px 0 10px 0;'>Changes aprovadas:</p>"
        f"{changes_html}"
        "</td></tr></table>"

        # FOOTER
        "<table width='100%' cellpadding='0' cellspacing='0' border='0'"
        "       bgcolor='#37B4C3' class='hdr-bg'"
        "       style='background-color:#37B4C3 !important;margin-top:40px;'>"
        "<tr><td align='center' valign='middle' bgcolor='#37B4C3' class='hdr-bg'"
        "    style='background-color:#37B4C3 !important;padding:0;margin:0;height:120px;'>"
        f"{footer_img}"
        "</td></tr></table>"

        "<p class='t-dark' style='font-size:16px;color:#222222 !important;mso-color-alt:#222222;margin:20px 40px 0 40px;'>Votos de um ótimo trabalho!</p>"
        "</body></html>"
    )


def montar_corpo_html(regiao: str, secroes_dados: dict[str, list[dict]]) -> str:
    data_formatada = datetime.now().strftime("%d/%m/%Y")
    titulo_regiao = "Ibéria" if regiao == "ib" else "South America - Brasil"

    html = (
        "<html><body style='font-family:\"EDP Preon\", Calibri, Arial, sans-serif;font-size:12pt;'>"
        "<table border='0' cellpadding='5' cellspacing='0' style='width:100%;margin-bottom:10px;'>"
        "<tr><td style='width:80%;text-align:center;vertical-align:middle;'>"
        f"<p style='color:#28FF52;font-family:\"EDP Preon\", Calibri, Arial, sans-serif;"
        f"font-size:13pt;font-weight:bold;margin:0;line-height:1.3;'>"
        f"Global Command Center EDP - {titulo_regiao}<br>Conf Call 09H</p>"
        "</td></tr></table>"
        f"<p style='color:#28FF52;font-weight:bold;margin:5px 0 10px 0;'>DIA: {data_formatada}</p>"
        "<hr style='border:0;border-top:1px solid #000;margin:10px 0;'>"
    )

    for titulo_secao, linhas in secroes_dados.items():
        html += (
            f"<p style='margin-top:20px;margin-bottom:8px;font-weight:bold;'>"
            f"<div style='background-color:#D9D9D9;padding:2px 4px;font-weight:bold;'>"
            f"{_esc(titulo_secao)}</div></p>"
        )
        if not linhas:
            html += "<p style='margin-left:15px;'>* Não foram identificadas situações.</p><br>"
        else:
            html += (
                "<br><table border='1' cellpadding='5' cellspacing='0' "
                "style='border-collapse:collapse;font-family:Calibri,Arial,sans-serif;"
                "font-size:9pt;width:100%;border:1px solid #000;'>"
            )
            cabecalhos = linhas[0].keys()
            html += "<tr style='background-color:#28FF52;color:#000000;font-weight:bold;'>"
            for cab in cabecalhos:
                html += f"<th style='padding:6px;text-align:left;border:1px solid #000;'>{_esc(cab)}</th>"
            html += "</tr>"
            for item in linhas:
                html += "<tr style='background-color:#FFFFFF;'>"
                for val in item.values():
                    html += f"<td style='padding:5px;border:1px solid #000;'>{_esc(val)}</td>"
                html += "</tr>"
            html += "</table><br>"

    html += "</body></html>"
    return html


# ------------------------------------------------------------------------------
# FUNÇÕES DE EMAIL
# ------------------------------------------------------------------------------

def gerar_email_p1_semanal(linhas: list[dict]) -> tuple[bool, str]:
    """Report P1 Semanal. Contactos lidos da sheet P1 SEMANA. Abre um rascunho no Outlook (não envia)."""
    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        mail = outlook.CreateItem(0)

        str_to, str_cc = _ler_contactos(
            "P1 SEMANA",
            fallback_to="GlobalCommandCenter@edp.com",
            fallback_cc="",
        )

        mail.To = str_to
        mail.CC = str_cc
        mail.Subject = "Analise de P1 semana anterior"
        mail.HTMLBody = montar_corpo_html_p1_semanal(linhas)
        mail.Display()
        return True, "E-mail do P1 Semanal gerado com sucesso no Outlook!"
    except Exception as e:  # noqa: BLE001
        logger.exception("Falha ao gerar e-mail P1 Semanal")
        return False, str(e)


def gerar_email_cab(tipo: str, texto_inicio: str, texto_fim: str, changes: list[dict]) -> tuple[bool, str]:
    """Report CAB — diário e FDS. Contactos lidos da sheet CAB. Abre um rascunho no Outlook (não envia)."""
    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        mail = outlook.CreateItem(0)

        agora = datetime.now()
        if tipo == "FDS":
            titulo = (f"CAB - Plano de Atividades entre: {agora.day} de "
                      f"{_nome_mes(agora.month)} a {texto_fim}")
        else:
            titulo = (f"CAB - Plano de Atividades: {agora.day} de "
                      f"{_nome_mes(agora.month)} a {texto_fim}")

        str_to, str_cc = _ler_contactos(
            "CAB",
            fallback_to="GlobalCommandCenter@edp.com",
            fallback_cc="",
        )

        mail.To = str_to
        mail.CC = str_cc
        mail.Subject = titulo
        mail.HTMLBody = montar_corpo_html_cab(tipo, texto_inicio, texto_fim, changes)
        mail.Display()
        return True, "E-mail CAB gerado com sucesso no Outlook!"
    except Exception as e:  # noqa: BLE001
        logger.exception("Falha ao gerar e-mail CAB")
        return False, str(e)


def gerar_email_outlook(regiao: str, tipo_report: str, dados_tabelas: dict[str, list[dict]]) -> tuple[bool, str]:
    """Report Ibéria / Brasil — diário e FDS. Contactos lidos da sheet correta. Abre um rascunho no Outlook (não envia)."""
    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        mail = outlook.CreateItem(0)

        agora = datetime.now().strftime("%Y-%m-%d")

        if regiao == "ib":
            str_to, str_cc = _ler_contactos(
                "CALL MATINAL PT",
                fallback_to="CentrodeComando@edp.pt",
                fallback_cc="GlobalCommandCenter@edp.com;CentrodeComando@edp.pt",
            )
            mail.Subject = f"REPORT - SUMÁRIO - Conference Call da manhã - Ibéria - {agora}"
        else:
            str_to, str_cc = _ler_contactos(
                "CALL MATINAL BR",
                fallback_to="GlobalCommandCenter@edp.com",
                fallback_cc="GlobalCommandCenter@edp.com",
            )
            mail.Subject = f"REPORT - SUMÁRIO - Conference Call da manhã - South America - {agora}"

        mail.To = str_to
        mail.CC = str_cc
        mail.HTMLBody = montar_corpo_html(regiao, dados_tabelas)
        mail.Display()
        return True, "E-mail gerado com sucesso no Outlook!"
    except Exception as e:  # noqa: BLE001
        logger.exception("Falha ao gerar e-mail Ibéria/Brasil")
        return False, str(e)
